from __future__ import annotations

import argparse
import csv
import io
import json
import math
import re
import unicodedata
import zipfile
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


ID_ALIASES = {
    "稿件id",
    "作品id",
    "笔记id",
    "内容id",
    "视频id",
    "稿件编号",
    "作品编号",
    "manuscriptid",
    "workid",
    "noteid",
    "contentid",
    "videoid",
    "aid",
    "avid",
}
TITLE_ALIASES = {
    "内容标题",
    "稿件标题",
    "作品标题",
    "笔记标题",
    "视频标题",
    "标题",
    "title",
}
AUTHOR_ALIASES = {
    "作者",
    "作者昵称",
    "发布账号",
    "博主",
    "博主昵称",
    "up主",
    "up主昵称",
    "author",
    "uploader",
    "nickname",
}
LINK_ALIASES = {
    "笔记链接",
    "作品链接",
    "内容链接",
    "视频链接",
    "链接",
    "url",
    "link",
}
DATE_ALIASES = {
    "发布时间",
    "发布日期",
    "投稿时间",
    "创建时间",
    "publishtime",
    "publishdate",
}
LINK_ID_PATTERNS = [
    re.compile(
        r"(?:xiaohongshu\.com/(?:explore|discovery/item)/)"
        r"([A-Za-z0-9_-]{6,128})(?=[/?#&\s]|$)",
        re.I,
    ),
    re.compile(
        r"(?:bilibili\.com/video/)(?:av)?"
        r"([A-Za-z0-9_-]{3,128})(?=[/?#&\s]|$)",
        re.I,
    ),
    re.compile(
        r"(?:[?&](?:id|note_id|work_id|content_id|video_id)=)"
        r"([A-Za-z0-9_-]{3,128})(?=[&#\s]|$)",
        re.I,
    ),
]
EXCEL_ERROR_PREFIXES = (
    "#REF!",
    "#DIV/0!",
    "#VALUE!",
    "#NAME?",
    "#N/A",
    "#NUM!",
    "#NULL!",
)


@dataclass
class Dataset:
    path: Path
    label: str
    sheet_name: str
    encoding: str | None
    delimiter: str | None
    header_row: int
    headers: list[str]
    rows: list[dict[str, Any]]
    fields: dict[str, int | None]


def normalized_text(value: Any) -> str:
    if value is None:
        return ""
    text = unicodedata.normalize("NFKC", str(value))
    text = "".join(
        char
        for char in text
        if unicodedata.category(char) not in {"Cc", "Cf", "Zl", "Zp"}
    )
    return re.sub(r"\s+", " ", text).strip()


def header_key(value: Any) -> str:
    return re.sub(r"[\s_\-./\\:：()（）\[\]【】]+", "", normalized_text(value)).casefold()


def make_unique_headers(values: Iterable[Any]) -> list[str]:
    counts: Counter[str] = Counter()
    result: list[str] = []
    for index, value in enumerate(values, start=1):
        base = normalized_text(value) or f"未命名字段_{index}"
        counts[base] += 1
        result.append(base if counts[base] == 1 else f"{base}_{counts[base]}")
    return result


def decode_text(path: Path) -> tuple[str, str]:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "gbk"):
        try:
            return raw.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    raise ValueError(f"无法识别文本编码：{path}")


def detect_delimiter(text: str, suffix: str) -> str:
    if suffix == ".tsv":
        return "\t"
    sample = text[:65536]
    try:
        return csv.Sniffer().sniff(sample, delimiters=",\t;|").delimiter
    except csv.Error:
        return ","


def detect_header_index(
    rows: list[tuple[Any, ...] | list[Any]],
    requested_id: str | None,
    explicit_row: int | None,
) -> int:
    if explicit_row is not None:
        if explicit_row < 1 or explicit_row > len(rows):
            raise ValueError(f"表头行超出范围：{explicit_row}")
        return explicit_row - 1

    targets = {header_key(requested_id)} if requested_id else ID_ALIASES
    for index, row in enumerate(rows[:30]):
        keys = {header_key(value) for value in row if normalized_text(value)}
        if keys & targets:
            return index
    raise ValueError(
        "前30行未识别到ID表头。请使用 --left-id/--right-id 和 "
        "--left-header-row/--right-header-row 显式指定。"
    )


def resolve_field(
    headers: list[str],
    requested: str | None,
    aliases: set[str],
    required: bool = False,
) -> int | None:
    keyed = [(index, header_key(header)) for index, header in enumerate(headers)]
    if requested:
        target = header_key(requested)
        for index, key in keyed:
            if key == target:
                return index
        raise ValueError(
            f"找不到字段“{requested}”。可用字段：{', '.join(headers)}"
        )
    for index, key in keyed:
        if key in aliases:
            return index
    if required:
        raise ValueError(f"无法自动识别ID字段。可用字段：{', '.join(headers)}")
    return None


def read_csv_dataset(
    path: Path,
    label: str,
    requested_fields: dict[str, str | None],
    header_row: int | None,
) -> Dataset:
    text, encoding = decode_text(path)
    delimiter = detect_delimiter(text, path.suffix.casefold())
    raw_rows = list(csv.reader(io.StringIO(text, newline=""), delimiter=delimiter))
    if not raw_rows:
        raise ValueError(f"空文件：{path}")
    header_index = detect_header_index(
        raw_rows, requested_fields["id"], header_row
    )
    headers = make_unique_headers(raw_rows[header_index])
    rows: list[dict[str, Any]] = []
    for source_row, raw in enumerate(
        raw_rows[header_index + 1 :], start=header_index + 2
    ):
        values = [
            raw[index] if index < len(raw) else None
            for index in range(len(headers))
        ]
        if any(normalized_text(value) for value in values):
            rows.append({"source_row": source_row, "values": values})
    return finalize_dataset(
        path,
        label,
        "CSV",
        encoding,
        delimiter,
        header_index + 1,
        headers,
        rows,
        requested_fields,
    )


def read_xlsx_dataset(
    path: Path,
    label: str,
    requested_fields: dict[str, str | None],
    sheet_name: str | None,
    header_row: int | None,
) -> Dataset:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if sheet_name:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(
                f"工作表“{sheet_name}”不存在。可用工作表："
                + ", ".join(workbook.sheetnames)
            )
        sheet = workbook[sheet_name]
    else:
        sheet = workbook.worksheets[0]
    if (
        sheet.max_row is None
        or sheet.max_column is None
        or (sheet.max_row <= 1 and sheet.max_column <= 1)
    ):
        sheet.reset_dimensions()
    raw_rows = list(sheet.iter_rows(values_only=True))
    if not raw_rows:
        raise ValueError(f"空工作表：{path} / {sheet.title}")
    header_index = detect_header_index(
        raw_rows, requested_fields["id"], header_row
    )
    headers = make_unique_headers(raw_rows[header_index])
    rows: list[dict[str, Any]] = []
    for source_row, raw in enumerate(
        raw_rows[header_index + 1 :], start=header_index + 2
    ):
        values = [
            raw[index] if index < len(raw) else None
            for index in range(len(headers))
        ]
        if any(normalized_text(value) for value in values):
            rows.append({"source_row": source_row, "values": values})
    return finalize_dataset(
        path,
        label,
        sheet.title,
        None,
        None,
        header_index + 1,
        headers,
        rows,
        requested_fields,
    )


def finalize_dataset(
    path: Path,
    label: str,
    sheet_name: str,
    encoding: str | None,
    delimiter: str | None,
    header_row: int,
    headers: list[str],
    rows: list[dict[str, Any]],
    requested_fields: dict[str, str | None],
) -> Dataset:
    fields = {
        "id": resolve_field(
            headers, requested_fields["id"], ID_ALIASES, required=True
        ),
        "title": resolve_field(
            headers, requested_fields["title"], TITLE_ALIASES
        ),
        "author": resolve_field(
            headers, requested_fields["author"], AUTHOR_ALIASES
        ),
        "link": resolve_field(
            headers, requested_fields["link"], LINK_ALIASES
        ),
        "date": resolve_field(
            headers, requested_fields["date"], DATE_ALIASES
        ),
    }
    return Dataset(
        path=path,
        label=label,
        sheet_name=sheet_name,
        encoding=encoding,
        delimiter=delimiter,
        header_row=header_row,
        headers=headers,
        rows=rows,
        fields=fields,
    )


def read_dataset(
    path: Path,
    label: str,
    requested_fields: dict[str, str | None],
    sheet_name: str | None,
    header_row: int | None,
) -> Dataset:
    suffix = path.suffix.casefold()
    if suffix in {".csv", ".tsv"}:
        return read_csv_dataset(path, label, requested_fields, header_row)
    if suffix in {".xlsx", ".xlsm"}:
        return read_xlsx_dataset(
            path, label, requested_fields, sheet_name, header_row
        )
    raise ValueError(
        f"不支持的文件类型：{path.suffix}。仅支持 CSV、TSV、XLSX、XLSM。"
    )


def normalize_token(token: str, case_sensitive: bool) -> str:
    return token if case_sensitive else token.casefold()


def extract_link_id(value: Any, case_sensitive: bool) -> str:
    text = normalized_text(value)
    for pattern in LINK_ID_PATTERNS:
        match = pattern.search(text)
        if match:
            return normalize_token(match.group(1), case_sensitive)
    return ""


def ids_are_comparable(left_id: str, right_id: str) -> bool:
    if left_id.isdecimal() != right_id.isdecimal():
        return False
    left_is_bv = left_id.casefold().startswith("bv")
    right_is_bv = right_id.casefold().startswith("bv")
    return left_is_bv == right_is_bv


def canonical_id(value: Any, case_sensitive: bool) -> tuple[str, bool]:
    if value is None or isinstance(value, bool):
        return "", False
    precision_risk = False
    if isinstance(value, int):
        return str(value), abs(value) >= 10**15
    if isinstance(value, float):
        if not math.isfinite(value) or not value.is_integer():
            return "", False
        precision_risk = abs(value) >= 10**15
        return format(value, ".0f"), precision_risk

    text = normalized_text(value)
    while text.startswith("'"):
        text = text[1:].lstrip()
    compact = re.sub(r"\s+", "", text)
    if not compact:
        return "", False

    link_id = extract_link_id(compact, case_sensitive)
    if link_id:
        return link_id, False

    if re.fullmatch(r"\+?\d+", compact):
        return compact.lstrip("+"), False
    decimal_match = re.fullmatch(r"\+?(\d+)\.0+", compact)
    if decimal_match:
        return decimal_match.group(1), False
    if re.fullmatch(r"[+-]?\d+(?:\.\d+)?[eE][+-]?\d+", compact):
        try:
            number = Decimal(compact)
            if number == number.to_integral_value() and number >= 0:
                return format(number.quantize(Decimal("1")), "f"), False
        except InvalidOperation:
            pass
    if re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_.:-]{5,127}", compact):
        return normalize_token(compact, case_sensitive), False
    return "", False


def parse_time(value: Any) -> float | None:
    if isinstance(value, datetime):
        return value.timestamp()
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day).timestamp()
    text = normalized_text(value).replace("/", "-")
    if not text:
        return None
    try:
        return datetime.fromisoformat(text).timestamp()
    except ValueError:
        return None


def field_value(dataset: Dataset, item: dict[str, Any], field: str) -> Any:
    index = dataset.fields[field]
    return item["values"][index] if index is not None else None


def prepare_records(dataset: Dataset, case_sensitive: bool) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    for item in dataset.rows:
        raw_id = field_value(dataset, item, "id")
        identifier, precision_risk = canonical_id(raw_id, case_sensitive)
        link_id = extract_link_id(
            field_value(dataset, item, "link"), case_sensitive
        )
        title = field_value(dataset, item, "title")
        author = field_value(dataset, item, "author")
        normalized_title = normalized_text(title).casefold()
        normalized_author = normalized_text(author).casefold()
        records.append(
            {
                "source_row": item["source_row"],
                "values": item["values"],
                "raw_id": "" if raw_id is None else str(raw_id),
                "trimmed_id": normalized_text(raw_id),
                "id": identifier,
                "precision_risk": precision_risk,
                "link_id": link_id,
                "link_mismatch": bool(
                    identifier
                    and link_id
                    and ids_are_comparable(identifier, link_id)
                    and identifier != link_id
                ),
                "title": title,
                "author": author,
                "title_author_key": (
                    (normalized_title, normalized_author)
                    if normalized_title and normalized_author
                    else None
                ),
                "publish_time": field_value(dataset, item, "date"),
            }
        )
    return records


def group_by_id(
    records: list[dict[str, Any]]
) -> dict[str, list[dict[str, Any]]]:
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for record in records:
        if record["id"]:
            grouped[record["id"]].append(record)
    return grouped


def pair_title_author(
    left_records: list[dict[str, Any]],
    right_records: list[dict[str, Any]],
) -> tuple[list[dict[str, Any]], dict[int, dict[str, Any]], dict[int, dict[str, Any]]]:
    left_groups: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    right_groups: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for record in left_records:
        if record["title_author_key"]:
            left_groups[record["title_author_key"]].append(record)
    for record in right_records:
        if record["title_author_key"]:
            right_groups[record["title_author_key"]].append(record)

    pairs: list[dict[str, Any]] = []
    left_map: dict[int, dict[str, Any]] = {}
    right_map: dict[int, dict[str, Any]] = {}
    for key in sorted(set(left_groups) & set(right_groups)):
        left = sorted(left_groups[key], key=lambda row: row["source_row"])
        remaining = sorted(
            right_groups[key], key=lambda row: row["source_row"]
        )
        for left_record in left:
            if not remaining:
                break
            left_time = parse_time(left_record["publish_time"])
            right_record = min(
                remaining,
                key=lambda row: (
                    left_time is None or parse_time(row["publish_time"]) is None,
                    abs(
                        (left_time or 0)
                        - (parse_time(row["publish_time"]) or 0)
                    ),
                    row["source_row"],
                ),
            )
            remaining.remove(right_record)
            pair = {
                "title": left_record["title"],
                "author": left_record["author"],
                "left_source_row": left_record["source_row"],
                "right_source_row": right_record["source_row"],
                "left_id": left_record["id"],
                "right_id": right_record["id"],
            }
            pairs.append(pair)
            left_map[left_record["source_row"]] = pair
            right_map[right_record["source_row"]] = pair
    return pairs, left_map, right_map


def duplicate_groups(records: list[dict[str, Any]]) -> int:
    return sum(1 for rows in group_by_id(records).values() if len(rows) > 1)


def date_range(records: list[dict[str, Any]]) -> tuple[str, str]:
    values: list[str] = []
    for record in records:
        value = record["publish_time"]
        if isinstance(value, (datetime, date)):
            values.append(value.isoformat()[:10])
            continue
        text = normalized_text(value)[:10]
        if re.fullmatch(r"\d{4}-\d{2}-\d{2}", text):
            values.append(text)
    return (min(values), max(values)) if values else ("", "")


def raw_set(records: list[dict[str, Any]], field: str) -> set[str]:
    return {record[field] for record in records if record[field]}


def analyze(
    left: Dataset,
    right: Dataset,
    left_records: list[dict[str, Any]],
    right_records: list[dict[str, Any]],
    skip_title_author: bool,
) -> dict[str, Any]:
    left_by_id = group_by_id(left_records)
    right_by_id = group_by_id(right_records)
    left_ids = set(left_by_id)
    right_ids = set(right_by_id)
    matched_ids = left_ids & right_ids
    left_only_ids = left_ids - right_ids
    right_only_ids = right_ids - left_ids

    matched_rows: list[list[Any]] = []
    for left_record in left_records:
        for right_record in right_by_id.get(left_record["id"], []):
            matched_rows.append(
                [
                    left_record["id"],
                    left_record["source_row"],
                    right_record["source_row"],
                ]
                + left_record["values"]
                + right_record["values"]
            )

    left_unmatched = [
        record
        for record in left_records
        if not record["id"] or record["id"] in left_only_ids
    ]
    right_unmatched = [
        record
        for record in right_records
        if not record["id"] or record["id"] in right_only_ids
    ]
    if skip_title_author:
        supplement_pairs, left_pair_map, right_pair_map = [], {}, {}
    else:
        supplement_pairs, left_pair_map, right_pair_map = pair_title_author(
            left_unmatched, right_unmatched
        )

    matched_headers = (
        ["匹配ID", "文件A源行号", "文件B源行号"]
        + [f"文件A_{header}" for header in left.headers]
        + [f"文件B_{header}" for header in right.headers]
    )
    unmatched_headers = [
        "来源行号",
        "匹配状态",
        "规范化ID",
        "标题作者补充匹配",
        "补充匹配对方行号",
        "补充匹配对方ID",
    ]

    left_unmatched_rows = []
    for record in left_unmatched:
        pair = left_pair_map.get(record["source_row"])
        left_unmatched_rows.append(
            [
                record["source_row"],
                "仅文件A存在" if record["id"] else "ID为空或格式异常",
                record["id"],
                "是" if pair else "否",
                pair["right_source_row"] if pair else None,
                pair["right_id"] if pair else None,
            ]
            + record["values"]
        )
    right_unmatched_rows = []
    for record in right_unmatched:
        pair = right_pair_map.get(record["source_row"])
        right_unmatched_rows.append(
            [
                record["source_row"],
                "仅文件B存在" if record["id"] else "ID为空或格式异常",
                record["id"],
                "是" if pair else "否",
                pair["left_source_row"] if pair else None,
                pair["left_id"] if pair else None,
            ]
            + record["values"]
        )

    if left_ids == right_ids:
        relation = "两个文件的有效规范化ID集合完全一致"
    elif right_ids < left_ids:
        relation = "文件B是文件A的严格子集"
    elif left_ids < right_ids:
        relation = "文件A是文件B的严格子集"
    elif matched_ids:
        relation = "两个文件部分重叠，互不包含"
    else:
        relation = "两个文件不存在共同有效规范化ID"

    left_links = raw_set(left_records, "link_id")
    right_links = raw_set(right_records, "link_id")
    summary = {
        "left_records": len(left_records),
        "right_records": len(right_records),
        "left_valid_id_records": sum(bool(row["id"]) for row in left_records),
        "right_valid_id_records": sum(bool(row["id"]) for row in right_records),
        "left_unique_ids": len(left_ids),
        "right_unique_ids": len(right_ids),
        "matched_unique_ids": len(matched_ids),
        "inner_join_rows": len(matched_rows),
        "left_unmatched_records": len(left_unmatched),
        "right_unmatched_records": len(right_unmatched),
        "left_only_unique_ids": len(left_only_ids),
        "right_only_unique_ids": len(right_only_ids),
        "left_invalid_id_records": sum(not row["id"] for row in left_records),
        "right_invalid_id_records": sum(not row["id"] for row in right_records),
        "left_duplicate_id_groups": duplicate_groups(left_records),
        "right_duplicate_id_groups": duplicate_groups(right_records),
        "left_match_rate": len(matched_ids) / len(left_ids) if left_ids else 0,
        "right_match_rate": len(matched_ids) / len(right_ids) if right_ids else 0,
        "left_is_subset_of_right": left_ids <= right_ids,
        "right_is_subset_of_left": right_ids <= left_ids,
        "relation": relation,
        "title_author_additional_pairs": len(supplement_pairs),
        "adjusted_left_unmatched_records": len(left_unmatched)
        - len(supplement_pairs),
        "adjusted_right_unmatched_records": len(right_unmatched)
        - len(supplement_pairs),
    }
    audit = {
        "raw_exact_matches": len(
            raw_set(left_records, "raw_id") & raw_set(right_records, "raw_id")
        ),
        "trimmed_matches": len(
            raw_set(left_records, "trimmed_id")
            & raw_set(right_records, "trimmed_id")
        ),
        "normalized_matches": len(matched_ids),
        "link_extraction_matches": len(left_links & right_links),
        "left_missing_link_ids": sum(
            bool(row["id"]) and not row["link_id"] for row in left_records
        ),
        "right_missing_link_ids": sum(
            bool(row["id"]) and not row["link_id"] for row in right_records
        ),
        "left_id_link_mismatches": sum(
            row["link_mismatch"] for row in left_records
        ),
        "right_id_link_mismatches": sum(
            row["link_mismatch"] for row in right_records
        ),
        "left_numeric_precision_risks": sum(
            row["precision_risk"] for row in left_records
        ),
        "right_numeric_precision_risks": sum(
            row["precision_risk"] for row in right_records
        ),
    }
    return {
        "summary": summary,
        "audit": audit,
        "matched_ids": matched_ids,
        "left_ids": left_ids,
        "right_ids": right_ids,
        "supplement_pairs": supplement_pairs,
        "sheets": [
            {
                "name": "匹配表_inner join",
                "headers": matched_headers,
                "rows": matched_rows,
                "theme": "matched",
                "freeze_columns": 3,
            },
            {
                "name": "不匹配详情_文件A",
                "headers": unmatched_headers + left.headers,
                "rows": left_unmatched_rows,
                "theme": "left",
                "freeze_columns": 6,
            },
            {
                "name": "不匹配详情_文件B",
                "headers": unmatched_headers + right.headers,
                "rows": right_unmatched_rows,
                "theme": "right",
                "freeze_columns": 6,
            },
        ],
        "date_ranges": {
            "left": date_range(left_records),
            "right": date_range(right_records),
        },
    }


def safe_excel_value(header: str, value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, datetime) and value.tzinfo is not None:
        value = value.replace(tzinfo=None)
    if isinstance(value, (datetime, date, int, float, bool)):
        if "id" not in header_key(header):
            return value
    if isinstance(value, (list, dict, tuple, set)):
        value = json.dumps(value, ensure_ascii=False)
    text = ILLEGAL_CHARACTERS_RE.sub("", str(value))
    return text[:32767]


def display_width(value: Any) -> int:
    text = "" if value is None else str(value)
    return sum(2 if unicodedata.east_asian_width(char) in {"W", "F"} else 1 for char in text)


def width_for_column(header: str, values: list[Any]) -> float:
    key = header_key(header)
    if "id" in key:
        return 28
    if "链接" in header or key in {"url", "link"}:
        return 38
    if "标题" in header:
        return 34
    if "正文" in header or "内容" == header or "摘要" in header:
        return 42
    if "时间" in header or "日期" in header:
        return 20
    sample_width = max(
        [display_width(header)] + [display_width(value) for value in values[:200]]
    )
    return min(max(sample_width + 2, 10), 24)


def append_typed_row(ws, headers: list[str], values: list[Any]) -> None:
    prepared = [
        safe_excel_value(headers[index], value)
        for index, value in enumerate(values)
    ]
    ws.append(prepared)
    for cell in ws[ws.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = "s"


def style_data_sheet(
    ws,
    headers: list[str],
    rows: list[list[Any]],
    theme: str,
    freeze_columns: int,
    table_index: int,
) -> None:
    colors = {
        "matched": ("2E7D32", "E8F5E9"),
        "left": ("9F1239", "FFF1F2"),
        "right": ("B45309", "FFF7ED"),
    }
    header_color, light_color = colors[theme]
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = f"{get_column_letter(freeze_columns + 1)}2"
    ws.auto_filter.ref = (
        f"A1:{get_column_letter(len(headers))}{max(len(rows) + 1, 1)}"
    )
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=header_color)
        cell.font = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
    ws.row_dimensions[1].height = 38

    for index, header in enumerate(headers, start=1):
        values = [row[index - 1] for row in rows if index <= len(row)]
        ws.column_dimensions[get_column_letter(index)].width = width_for_column(
            header, values
        )
        if rows and "id" in header_key(header):
            for cell in ws.iter_cols(
                min_col=index,
                max_col=index,
                min_row=2,
                max_row=len(rows) + 1,
            ):
                for item in cell:
                    item.number_format = "0"

    if rows:
        ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
        table = Table(displayName=f"ManuscriptMatch{table_index}", ref=ref)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(table)

    if theme in {"left", "right"} and rows:
        for row_index in range(2, len(rows) + 2):
            ws.cell(row_index, 2).fill = PatternFill(
                "solid", fgColor=light_color
            )
            ws.cell(row_index, 2).font = Font(
                name="Aptos", size=10, bold=True, color=header_color
            )
            if ws.cell(row_index, 4).value == "是":
                for column in range(4, 7):
                    ws.cell(row_index, column).fill = PatternFill(
                        "solid", fgColor="E0F2FE"
                    )
                    ws.cell(row_index, column).font = Font(
                        name="Aptos", size=10, bold=True, color="075985"
                    )


def add_analysis_sheet(
    workbook: Workbook,
    left: Dataset,
    right: Dataset,
    result: dict[str, Any],
) -> None:
    ws = workbook.create_sheet("分析汇总")
    ws.sheet_view.showGridLines = False
    summary = result["summary"]
    audit = result["audit"]
    left_start, left_end = result["date_ranges"]["left"]
    right_start, right_end = result["date_ranges"]["right"]

    ws.merge_cells("A1:H1")
    ws["A1"] = "稿件ID匹配与差异分析汇总"
    ws["A1"].fill = PatternFill("solid", fgColor="17365D")
    ws["A1"].font = Font(
        name="Aptos Display", size=18, bold=True, color="FFFFFF"
    )
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 38
    ws.merge_cells("A2:H2")
    ws["A2"] = (
        f"文件A：{left.path}    文件B：{right.path}    "
        f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    )
    ws["A2"].font = Font(name="Aptos", size=9, color="475569")
    ws["A2"].alignment = Alignment(vertical="center")
    ws.row_dimensions[2].height = 24

    rows = [
        ("匹配概览", None, None, None),
        ("指标", left.label, right.label, "说明"),
        ("源数据记录数", summary["left_records"], summary["right_records"], "不含表头"),
        ("有效ID记录数", summary["left_valid_id_records"], summary["right_valid_id_records"], "规范化后非空"),
        ("唯一有效ID数", summary["left_unique_ids"], summary["right_unique_ids"], "去重后"),
        ("匹配唯一ID数", summary["matched_unique_ids"], summary["matched_unique_ids"], "ID交集"),
        ("inner join行数", summary["inner_join_rows"], summary["inner_join_rows"], "重复ID采用笛卡尔连接"),
        ("不匹配记录数", summary["left_unmatched_records"], summary["right_unmatched_records"], "包含无效ID"),
        ("仅本文件唯一ID数", summary["left_only_unique_ids"], summary["right_only_unique_ids"], "方向性差异"),
        ("匹配率/覆盖率", summary["left_match_rate"], summary["right_match_rate"], "匹配唯一ID / 本文件唯一有效ID"),
        ("重复ID组数", summary["left_duplicate_id_groups"], summary["right_duplicate_id_groups"], "同一规范化ID出现多次"),
        ("无效ID记录数", summary["left_invalid_id_records"], summary["right_invalid_id_records"], "空值或无法规范化"),
        ("有效ID集合包含关系", "是" if summary["left_is_subset_of_right"] else "否", "是" if summary["right_is_subset_of_left"] else "否", "本文件是否为对方子集"),
        ("发布日期范围", f"{left_start} 至 {left_end}" if left_start else "", f"{right_start} 至 {right_end}" if right_start else "", "自动识别发布时间字段"),
        ("格式与链接审计", None, None, None),
        ("指标", left.label, right.label, "说明"),
        ("原始ID集合交集", audit["raw_exact_matches"], audit["raw_exact_matches"], "未清洗文本"),
        ("去空白ID集合交集", audit["trimmed_matches"], audit["trimmed_matches"], "NFKC与空白折叠"),
        ("规范化ID集合交集", audit["normalized_matches"], audit["normalized_matches"], "最终精确匹配口径"),
        ("链接提取ID集合交集", audit["link_extraction_matches"], audit["link_extraction_matches"], "仅审计，不补写ID"),
        ("有ID但链接未提取", audit["left_missing_link_ids"], audit["right_missing_link_ids"], "未识别链接或链接为空"),
        ("ID与链接ID不一致", audit["left_id_link_mismatches"], audit["right_id_link_mismatches"], "两者均可提取时比较"),
        ("Excel长数字精度风险", audit["left_numeric_precision_risks"], audit["right_numeric_precision_risks"], "数值型且为16位及以上整数"),
        ("标题与作者补充判断", None, None, None),
        ("指标", left.label, right.label, "说明"),
        ("标题+作者补充匹配对", summary["title_author_additional_pairs"], summary["title_author_additional_pairs"], "仅在ID不匹配记录中一对一配对"),
        ("补充判断后仍不匹配", summary["adjusted_left_unmatched_records"], summary["adjusted_right_unmatched_records"], "不改变ID精确匹配口径"),
        ("最终结论", summary["relation"], summary["relation"], "按唯一有效规范化ID集合判断"),
    ]

    start_row = 4
    for offset, row in enumerate(rows):
        row_index = start_row + offset
        if row[1:] == (None, None, None):
            ws.merge_cells(
                start_row=row_index,
                start_column=1,
                end_row=row_index,
                end_column=4,
            )
            ws.cell(row_index, 1, row[0])
            ws.cell(row_index, 1).fill = PatternFill(
                "solid", fgColor="1F4E78"
            )
            ws.cell(row_index, 1).font = Font(
                name="Aptos", size=11, bold=True, color="FFFFFF"
            )
            ws.row_dimensions[row_index].height = 26
            continue
        for column, value in enumerate(row, start=1):
            cell = ws.cell(row_index, column, value)
            cell.font = Font(name="Aptos", size=10, color="1F2937")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = Border(
                bottom=Side(style="thin", color="E2E8F0")
            )
        if row[0] == "指标":
            for column in range(1, 5):
                cell = ws.cell(row_index, column)
                cell.fill = PatternFill("solid", fgColor="334155")
                cell.font = Font(
                    name="Aptos", size=10, bold=True, color="FFFFFF"
                )
                cell.alignment = Alignment(
                    horizontal="center", vertical="center", wrap_text=True
                )
            ws.row_dimensions[row_index].height = 28
        elif row[0] == "匹配率/覆盖率":
            ws.cell(row_index, 2).number_format = "0.00%"
            ws.cell(row_index, 3).number_format = "0.00%"

    ws.freeze_panes = "A4"
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 32
    ws.column_dimensions["C"].width = 32
    ws.column_dimensions["D"].width = 38
    for column in "EFGH":
        ws.column_dimensions[column].width = 4


def write_workbook(
    output: Path,
    left: Dataset,
    right: Dataset,
    result: dict[str, Any],
) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for table_index, sheet_data in enumerate(result["sheets"], start=1):
        ws = workbook.create_sheet(sheet_data["name"])
        append_typed_row(ws, sheet_data["headers"], sheet_data["headers"])
        for row in sheet_data["rows"]:
            append_typed_row(ws, sheet_data["headers"], row)
        style_data_sheet(
            ws,
            sheet_data["headers"],
            sheet_data["rows"],
            sheet_data["theme"],
            sheet_data["freeze_columns"],
            table_index,
        )
    add_analysis_sheet(workbook, left, right, result)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def validate_output(output: Path, result: dict[str, Any]) -> dict[str, Any]:
    with zipfile.ZipFile(output) as archive:
        bad_member = archive.testzip()
        if bad_member is not None:
            raise AssertionError(f"OOXML压缩包损坏：{bad_member}")

    workbook = load_workbook(output, read_only=False, data_only=False)
    expected_names = [
        "匹配表_inner join",
        "不匹配详情_文件A",
        "不匹配详情_文件B",
        "分析汇总",
    ]
    if workbook.sheetnames != expected_names:
        raise AssertionError(
            f"工作表顺序错误：{workbook.sheetnames}"
        )
    for sheet_data in result["sheets"]:
        ws = workbook[sheet_data["name"]]
        expected_rows = len(sheet_data["rows"]) + 1
        if ws.max_row != expected_rows:
            raise AssertionError(
                f"{ws.title} 行数错误：{ws.max_row} != {expected_rows}"
            )
        expected_columns = len(sheet_data["headers"])
        if ws.max_column != expected_columns:
            raise AssertionError(
                f"{ws.title} 列数错误：{ws.max_column} != {expected_columns}"
            )

    matched_ws = workbook["匹配表_inner join"]
    workbook_matched_ids = {
        matched_ws.cell(row, 1).value
        for row in range(2, matched_ws.max_row + 1)
    }
    if workbook_matched_ids != result["matched_ids"]:
        raise AssertionError("匹配表ID集合与计算结果不一致")
    for row in range(2, matched_ws.max_row + 1):
        value = matched_ws.cell(row, 1).value
        if not isinstance(value, str):
            raise AssertionError(f"匹配ID未按文本存储：A{row}")

    unmatched_checks = [
        ("不匹配详情_文件A", result["left_ids"] - result["matched_ids"]),
        ("不匹配详情_文件B", result["right_ids"] - result["matched_ids"]),
    ]
    for sheet_name, expected_ids in unmatched_checks:
        ws = workbook[sheet_name]
        actual_ids = {
            ws.cell(row, 3).value
            for row in range(2, ws.max_row + 1)
            if ws.cell(row, 3).value
        }
        if actual_ids != expected_ids:
            raise AssertionError(f"{sheet_name} ID集合与计算结果不一致")

    for sheet_data in result["sheets"]:
        ws = workbook[sheet_data["name"]]
        for column, header in enumerate(sheet_data["headers"], start=1):
            if "id" not in header_key(header):
                continue
            for row in range(2, ws.max_row + 1):
                value = ws.cell(row, column).value
                if value is not None and not isinstance(value, str):
                    raise AssertionError(
                        f"ID列未按文本存储：{ws.title}!{get_column_letter(column)}{row}"
                    )

    for ws in workbook.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "e":
                    raise AssertionError(
                        f"发现Excel错误值：{ws.title}!{cell.coordinate}"
                    )
                if isinstance(cell.value, str) and cell.value.startswith(
                    EXCEL_ERROR_PREFIXES
                ):
                    raise AssertionError(
                        f"发现错误文本：{ws.title}!{cell.coordinate}"
                    )
    return {
        "zip_integrity": "passed",
        "sheet_order": "passed",
        "shape_checks": "passed",
        "id_text_checks": "passed",
        "error_scan": "passed",
    }


def requested_fields(args: argparse.Namespace, side: str) -> dict[str, str | None]:
    return {
        name: getattr(args, f"{side}_{name}")
        for name in ("id", "title", "author", "link", "date")
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "按规范化稿件/作品/笔记ID匹配两个CSV/XLSX文件，"
            "输出inner join、双向差异和分析汇总。"
        )
    )
    parser.add_argument("left_file", type=Path, help="文件A路径")
    parser.add_argument("right_file", type=Path, help="文件B路径")
    parser.add_argument("--output", type=Path, help="输出XLSX路径")
    parser.add_argument("--json-report", type=Path, help="可选JSON结果路径")
    parser.add_argument("--left-sheet", help="文件A工作表名称")
    parser.add_argument("--right-sheet", help="文件B工作表名称")
    parser.add_argument("--left-header-row", type=int, help="文件A表头行，1起始")
    parser.add_argument("--right-header-row", type=int, help="文件B表头行，1起始")
    parser.add_argument("--left-label", default="文件A", help="文件A显示名称")
    parser.add_argument("--right-label", default="文件B", help="文件B显示名称")
    for side in ("left", "right"):
        for field, help_text in (
            ("id", "ID字段"),
            ("title", "标题字段"),
            ("author", "作者字段"),
            ("link", "链接字段"),
            ("date", "发布时间字段"),
        ):
            parser.add_argument(
                f"--{side}-{field}",
                dest=f"{side}_{field}",
                help=f"{'文件A' if side == 'left' else '文件B'}{help_text}",
            )
    parser.add_argument(
        "--id-case-sensitive",
        action="store_true",
        help="非数字ID区分大小写；默认不区分",
    )
    parser.add_argument(
        "--skip-title-author",
        action="store_true",
        help="不执行标题+作者补充匹配",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    left_path = args.left_file.resolve()
    right_path = args.right_file.resolve()
    if not left_path.is_file():
        raise FileNotFoundError(f"文件A不存在：{left_path}")
    if not right_path.is_file():
        raise FileNotFoundError(f"文件B不存在：{right_path}")
    output = (
        args.output.resolve()
        if args.output
        else (
            Path.cwd()
            / f"{left_path.stem}_{right_path.stem}_稿件ID匹配结果.xlsx"
        ).resolve()
    )
    if output in {left_path, right_path}:
        raise ValueError("输出路径不得覆盖源文件")

    left = read_dataset(
        left_path,
        args.left_label,
        requested_fields(args, "left"),
        args.left_sheet,
        args.left_header_row,
    )
    right = read_dataset(
        right_path,
        args.right_label,
        requested_fields(args, "right"),
        args.right_sheet,
        args.right_header_row,
    )
    left_records = prepare_records(left, args.id_case_sensitive)
    right_records = prepare_records(right, args.id_case_sensitive)
    result = analyze(
        left,
        right,
        left_records,
        right_records,
        args.skip_title_author,
    )
    write_workbook(output, left, right, result)
    validation = validate_output(output, result)

    report = {
        "output": str(output),
        "sources": {
            "left": str(left.path),
            "right": str(right.path),
        },
        "source_metadata": {
            "left_sheet": left.sheet_name,
            "right_sheet": right.sheet_name,
            "left_header_row": left.header_row,
            "right_header_row": right.header_row,
            "left_encoding": left.encoding,
            "right_encoding": right.encoding,
        },
        "field_mapping": {
            "left": {
                name: left.headers[index] if index is not None else None
                for name, index in left.fields.items()
            },
            "right": {
                name: right.headers[index] if index is not None else None
                for name, index in right.fields.items()
            },
        },
        "summary": result["summary"],
        "normalization_and_link_audit": result["audit"],
        "validation": validation,
    }
    if args.json_report:
        json_path = args.json_report.resolve()
        json_path.parent.mkdir(parents=True, exist_ok=True)
        json_path.write_text(
            json.dumps(report, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
