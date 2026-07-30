from __future__ import annotations

import csv
import json
import os
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook


SCRIPT = Path(__file__).with_name("match_manuscript_ids.py")


def main() -> int:
    with tempfile.TemporaryDirectory(prefix="match-manuscript-ids-") as temp:
        root = Path(temp)
        left = root / "left.csv"
        right = root / "right.xlsx"
        output = root / "result.xlsx"
        report = root / "result.json"

        with left.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(
                ["稿件ID", "内容标题", "作者昵称", "笔记链接", "发布时间"]
            )
            writer.writerows(
                [
                    [" 00123 ", "标题A", "作者A", "https://www.bilibili.com/video/av00123", "2026-01-01"],
                    ["456.0", "标题B", "作者B", "https://www.bilibili.com/video/av456", "2026-01-02"],
                    ["7.89E+2", "标题C", "作者C", "https://www.bilibili.com/video/av789", "2026-01-03"],
                    ["900", "标题D", "作者D", "https://www.bilibili.com/video/av900", "2026-01-04"],
                    ["1000", "相同标题", "相同作者", "https://www.bilibili.com/video/av1000", "2026-01-05"],
                    ["", "无效ID", "作者E", "", "2026-01-06"],
                ]
            )

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "详情"
        sheet.append(["测试报告"])
        sheet.append([])
        sheet.append(["作品ID", "内容标题", "作者", "链接", "发布时间"])
        sheet.append(["00123", "标题A", "作者A", "https://www.bilibili.com/video/av00123", "2026-01-01"])
        sheet.append([456, "标题B", "作者B", "https://www.bilibili.com/video/av456", "2026-01-02"])
        sheet.append([789, "标题C", "作者C", "https://www.bilibili.com/video/av789", "2026-01-03"])
        sheet.append([900, "标题D", "作者D", "https://www.bilibili.com/video/av900", "2026-01-04"])
        sheet.append([2000, "相同标题", "相同作者", "https://www.bilibili.com/video/av2000", "2026-01-05"])
        workbook.save(right)

        completed = subprocess.run(
            [
                sys.executable,
                str(SCRIPT),
                str(left),
                str(right),
                "--left-id",
                "稿件ID",
                "--right-id",
                "作品ID",
                "--output",
                str(output),
                "--json-report",
                str(report),
            ],
            check=True,
            capture_output=True,
            text=True,
            encoding="utf-8",
            env={**os.environ, "PYTHONIOENCODING": "utf-8"},
        )
        payload = json.loads(report.read_text(encoding="utf-8"))
        summary = payload["summary"]
        assert summary["matched_unique_ids"] == 4
        assert summary["inner_join_rows"] == 4
        assert summary["left_unmatched_records"] == 2
        assert summary["right_unmatched_records"] == 1
        assert summary["title_author_additional_pairs"] == 1
        assert summary["relation"] == "两个文件部分重叠，互不包含"
        assert all(
            value == "passed" for value in payload["validation"].values()
        )

        result_workbook = load_workbook(output, data_only=False)
        assert result_workbook.sheetnames == [
            "匹配表_inner join",
            "不匹配详情_文件A",
            "不匹配详情_文件B",
            "分析汇总",
        ]
        matched = result_workbook["匹配表_inner join"]
        ids = {matched.cell(row, 1).value for row in range(2, 6)}
        assert ids == {"00123", "456", "789", "900"}
        assert all(isinstance(value, str) for value in ids)
        assert output.is_file()
        print(
            json.dumps(
                {
                    "self_test": "passed",
                    "matched_ids": sorted(ids),
                    "validation": payload["validation"],
                    "stdout_bytes": len(completed.stdout.encode("utf-8")),
                },
                ensure_ascii=False,
                indent=2,
            )
        )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
