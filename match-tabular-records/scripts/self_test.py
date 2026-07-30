from __future__ import annotations

import csv
import json
import os
import subprocess
import sys
import tempfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook


SCRIPT = Path(__file__).with_name("match_tabular_records.py")


def run_match(arguments: list[str], report: Path) -> dict:
    completed = subprocess.run(
        [sys.executable, str(SCRIPT), *arguments],
        check=True,
        capture_output=True,
        text=True,
        encoding="utf-8",
        env={**os.environ, "PYTHONIOENCODING": "utf-8"},
    )
    payload = json.loads(report.read_text(encoding="utf-8"))
    assert all(
        value == "passed" for value in payload["validation"].values()
    )
    payload["_stdout_bytes"] = len(completed.stdout.encode("utf-8"))
    return payload


def single_key_test(root: Path) -> dict:
    left = root / "single-left.csv"
    right = root / "single-right.xlsx"
    output = root / "single-result.xlsx"
    report = root / "single-result.json"

    with left.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(
            ["稿件ID", "内容标题", "作者昵称", "发布时间"]
        )
        writer.writerows(
            [
                [" 00123 ", "标题A", "作者A", "2026-01-01"],
                ["456.0", "标题B", "作者B", "2026-01-02"],
                ["7.89E+2", "标题C", "作者C", "2026-01-03"],
                ["900", "标题D", "作者D", "2026-01-04"],
                ["1000", "相同标题", "相同作者", "2026-01-05"],
                ["", "无效ID", "作者E", "2026-01-06"],
            ]
        )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "详情"
    sheet.append(["测试报告"])
    sheet.append([])
    sheet.append(["作品ID", "内容标题", "作者", "发布时间"])
    sheet.append(["00123", "标题A", "作者A", "2026-01-01"])
    sheet.append([456, "标题B", "作者B", "2026-01-02"])
    sheet.append([789, "标题C", "作者C", "2026-01-03"])
    sheet.append([900, "标题D", "作者D", "2026-01-04"])
    sheet.append([2000, "相同标题", "相同作者", "2026-01-05"])
    workbook.save(right)

    payload = run_match(
        [
            str(left),
            str(right),
            "--left-key",
            "稿件ID",
            "--right-key",
            "作品ID",
            "--key-normalizer",
            "id",
            "--left-secondary-key",
            "内容标题",
            "--right-secondary-key",
            "内容标题",
            "--left-secondary-key",
            "作者昵称",
            "--right-secondary-key",
            "作者",
            "--secondary-normalizer",
            "text",
            "--left-date",
            "发布时间",
            "--right-date",
            "发布时间",
            "--output",
            str(output),
            "--json-report",
            str(report),
        ],
        report,
    )
    summary = payload["summary"]
    assert summary["matched_unique_keys"] == 4
    assert summary["inner_join_rows"] == 4
    assert summary["left_unmatched_records"] == 2
    assert summary["right_unmatched_records"] == 1
    assert summary["secondary_additional_pairs"] == 1
    assert summary["relation"] == "两个文件部分重叠，互不包含"

    result = load_workbook(output, data_only=False)
    assert result.sheetnames == [
        "匹配表_inner join",
        "不匹配详情_文件A",
        "不匹配详情_文件B",
        "分析汇总",
    ]
    matched = result["匹配表_inner join"]
    keys = {matched.cell(row, 1).value for row in range(2, 6)}
    assert keys == {"00123", "456", "789", "900"}
    assert all(isinstance(value, str) for value in keys)
    return {
        "matched_keys": sorted(keys),
        "secondary_pairs": summary["secondary_additional_pairs"],
        "stdout_bytes": payload["_stdout_bytes"],
    }


def composite_key_test(root: Path) -> dict:
    left = root / "composite-left.csv"
    right = root / "composite-right.xlsx"
    output = root / "composite-result.xlsx"
    report = root / "composite-result.json"

    with left.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["标题", "金额", "日期", "备注"])
        writer.writerows(
            [
                [" Alpha   Story ", "1,000.00", "2026/01/01", "A"],
                ["Beta", "20.0", "2026-01-02", "B"],
                ["Gamma", "3E1", "2026年1月3日", "C"],
                ["Only Left", "40", "2026-01-04", "D"],
                ["", "50", "2026-01-05", "invalid"],
            ]
        )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "业务数据"
    sheet.append(["内容标题", "金额值", "业务日期", "说明"])
    sheet.append(["alpha story", 1000, datetime(2026, 1, 1, 8, 0), "A"])
    sheet.append(["Ｂｅｔａ", 20, datetime(2026, 1, 2), "B"])
    sheet.append(["GAMMA", 30, datetime(2026, 1, 3, 23, 59), "C"])
    sheet.append(["Only Right", 60, datetime(2026, 1, 6), "E"])
    workbook.save(right)

    payload = run_match(
        [
            str(left),
            str(right),
            "--left-key",
            "标题",
            "--right-key",
            "内容标题",
            "--left-key",
            "金额",
            "--right-key",
            "金额值",
            "--left-key",
            "日期",
            "--right-key",
            "业务日期",
            "--key-normalizer",
            "text",
            "--key-normalizer",
            "number",
            "--key-normalizer",
            "date",
            "--output",
            str(output),
            "--json-report",
            str(report),
        ],
        report,
    )
    summary = payload["summary"]
    assert summary["matched_unique_keys"] == 3
    assert summary["inner_join_rows"] == 3
    assert summary["left_unmatched_records"] == 2
    assert summary["right_unmatched_records"] == 1
    assert summary["left_invalid_key_records"] == 1
    assert summary["secondary_additional_pairs"] == 0

    result = load_workbook(output, data_only=False)
    matched = result["匹配表_inner join"]
    keys = {
        tuple(matched.cell(row, column).value for column in range(1, 4))
        for row in range(2, 5)
    }
    assert keys == {
        ("alpha story", "1000", "2026-01-01"),
        ("beta", "20", "2026-01-02"),
        ("gamma", "30", "2026-01-03"),
    }
    return {
        "matched_keys": [list(key) for key in sorted(keys)],
        "invalid_left_keys": summary["left_invalid_key_records"],
        "stdout_bytes": payload["_stdout_bytes"],
    }


def main() -> int:
    with tempfile.TemporaryDirectory(
        prefix="match-tabular-records-"
    ) as temp:
        root = Path(temp)
        result = {
            "self_test": "passed",
            "single_key": single_key_test(root),
            "composite_key": composite_key_test(root),
        }
        print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
