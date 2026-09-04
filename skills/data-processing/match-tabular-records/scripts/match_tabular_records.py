from __future__ import annotations

import argparse
import csv
import io
import json
import math
import re
import unicodedata
import zipfile
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


AUTO_ID_ALIASES = {
    "id",
    "稿件id",
    "作品id",
    "笔记id",
    "内容id",
    "视频id",
    "记录id",
    "订单id",
    "稿件编号",
    "作品编号",
    "记录编号",
    "订单编号",
    "manuscriptid",
    "workid",
    "noteid",
    "contentid",
    "videoid",
    "recordid",
    "orderid",
    "aid",
    "avid",
}
DATE_ALIASES = {
    "发布时间",
    "发布日期",
    "投稿时间",
    "创建时间",
    "更新时间",
    "日期",
    "时间",
    "date",
    "datetime",
    "publishtime",
    "publishdate",
    "createdat",
    "updatedat",
}
NORMALIZER_CHOICES = (
    "auto",
    "id",
    "text",
    "number",
    "date",
    "datetime",
    "exact",
)
LINK_ID_PATTERNS = [
    re.compile(
        r"(?:xiaohongshu\.com/(?:explore|discovery/item)/)"
        r"([A-Za-z0-9_-]{6,256})(?=[/?#&\s]|$)",
        re.I,
    ),
    re.compile(
        r"(?:bilibili\.com/video/)(?:av)?"
        r"([A-Za-z0-9_-]{3,256})(?=[/?#&\s]|$)",
        re.I,
    ),
    re.compile(
        r"(?:[?&](?:id|note_id|work_id|content_id|video_id|record_id)=)"
        r"([^&#\s]{1,256})(?=[&#\s]|$)",
        re.I,
    ),
]
EXCEL_ERROR_PREFIXES = (
    "#REF!",
    "#DIV/0!",
    "#VALUE!",
    "#NAME?",
    "#N/A",
    "#NUM!",
    "#NULL!",
)


@dataclass
class Dataset:
    path: Path
    label: str
    sheet_name: str
    encoding: str | None
    delimiter: str | None
    header_row: int
    headers: list[str]
    rows: list[dict[str, Any]]
    key_fields: list[int]
    secondary_fields: list[int]
    date_field: int | None


def scalar_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def normalized_text(value: Any, case_sensitive: bool = False) -> str:
    text = unicodedata.normalize("NFKC", scalar_text(value))
    text = "".join(
        " "
        if unicodedata.category(char) in {"Cc", "Cf", "Zl", "Zp"}
        else char
        for char in text
    )
    text = re.sub(r"\s+", " ", text).strip()
    return text if case_sensitive else text.casefold()


def header_key(value: Any) -> str:
    return re.sub(
        r"[\s_\-./\\:：()（）\[\]【】]+",
        "",
        normalized_text(value),
    )


def make_unique_headers(values: Iterable[Any]) -> list[str]:
    counts: Counter[str] = Counter()
    result: list[str] = []
    for index, value in enumerate(values, start=1):
        base = normalized_text(value, case_sensitive=True) or f"未命名字段_{index}"
        counts[base] += 1
        result.append(base if counts[base] == 1 else f"{base}_{counts[base]}")
    return result


def decode_text(path: Path) -> tuple[str, str]:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "gbk"):
        try:
            return raw.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    raise ValueError(f"无法识别文本编码：{path}")


def detect_delimiter(text: str, suffix: str) -> str:
    if suffix == ".tsv":
        return "\t"
    try:
        return csv.Sniffer().sniff(
            text[:65536], delimiters=",\t;|"
        ).delimiter
    except csv.Error:
        return ","


def detect_header_index(
    rows: list[tuple[Any, ...] | list[Any]],
    requested_keys: list[str],
    explicit_row: int | None,
) -> int:
    if explicit_row is not None:
        if explicit_row < 1 or explicit_row > len(rows):
            raise ValueError(f"表头行超出范围：{explicit_row}")
        return explicit_row - 1

    requested = {header_key(name) for name in requested_keys}
    for index, row in enumerate(rows[:30]):
        keys = {header_key(value) for value in row if scalar_text(value)}
        if requested and requested <= keys:
            return index
        if not requested and keys & AUTO_ID_ALIASES:
            return index
    if requested:
        detail = "、".join(requested_keys)
        raise ValueError(
            f"前30行未找到全部匹配字段：{detail}。请显式指定表头行。"
        )
    raise ValueError(
        "前30行未识别到默认ID字段。请使用 --left-key/--right-key "
        "显式指定一个或多个匹配字段。"
    )


def resolve_field(headers: list[str], requested: str) -> int:
    target = header_key(requested)
    matches = [
        index
        for index, header in enumerate(headers)
        if header_key(header) == target
    ]
    if len(matches) == 1:
        return matches[0]
    if not matches:
        raise ValueError(
            f"找不到字段“{requested}”。可用字段：{', '.join(headers)}"
        )
    raise ValueError(f"字段“{requested}”存在多个同名候选，请先明确表头。")


def resolve_auto_id_field(headers: list[str]) -> int:
    matches = [
        index
        for index, header in enumerate(headers)
        if header_key(header) in AUTO_ID_ALIASES
    ]
    if len(matches) == 1:
        return matches[0]
    if not matches:
        raise ValueError(
            "无法自动识别单一ID字段。请使用 --left-key/--right-key。"
        )
    names = ", ".join(headers[index] for index in matches)
    raise ValueError(f"发现多个ID候选字段：{names}。请显式指定匹配字段。")


def resolve_optional_date_field(
    headers: list[str], requested: str | None
) -> int | None:
    if requested:
        return resolve_field(headers, requested)
    for index, header in enumerate(headers):
        if header_key(header) in DATE_ALIASES:
            return index
    return None


def finalize_dataset(
    *,
    path: Path,
    label: str,
    sheet_name: str,
    encoding: str | None,
    delimiter: str | None,
    header_row: int,
    headers: list[str],
    rows: list[dict[str, Any]],
    requested_keys: list[str],
    requested_secondary: list[str],
    requested_date: str | None,
) -> Dataset:
    key_fields = (
        [resolve_field(headers, name) for name in requested_keys]
        if requested_keys
        else [resolve_auto_id_field(headers)]
    )
    secondary_fields = [
        resolve_field(headers, name) for name in requested_secondary
    ]
    return Dataset(
        path=path,
        label=label,
        sheet_name=sheet_name,
        encoding=encoding,
        delimiter=delimiter,
        header_row=header_row,
        headers=headers,
        rows=rows,
        key_fields=key_fields,
        secondary_fields=secondary_fields,
        date_field=resolve_optional_date_field(headers, requested_date),
    )


def read_csv_dataset(
    path: Path,
    label: str,
    requested_keys: list[str],
    requested_secondary: list[str],
    requested_date: str | None,
    header_row: int | None,
) -> Dataset:
    text, encoding = decode_text(path)
    delimiter = detect_delimiter(text, path.suffix.casefold())
    raw_rows = list(csv.reader(io.StringIO(text, newline=""), delimiter=delimiter))
    if not raw_rows:
        raise ValueError(f"空文件：{path}")
    header_index = detect_header_index(raw_rows, requested_keys, header_row)
    headers = make_unique_headers(raw_rows[header_index])
    rows: list[dict[str, Any]] = []
    for source_row, raw in enumerate(
        raw_rows[header_index + 1 :], start=header_index + 2
    ):
        values = [
            raw[index] if index < len(raw) else None
            for index in range(len(headers))
        ]
        if any(normalized_text(value) for value in values):
            rows.append({"source_row": source_row, "values": values})
    return finalize_dataset(
        path=path,
        label=label,
        sheet_name="CSV",
        encoding=encoding,
        delimiter=delimiter,
        header_row=header_index + 1,
        headers=headers,
        rows=rows,
        requested_keys=requested_keys,
        requested_secondary=requested_secondary,
        requested_date=requested_date,
    )


def read_xlsx_dataset(
    path: Path,
    label: str,
    requested_keys: list[str],
    requested_secondary: list[str],
    requested_date: str | None,
    sheet_name: str | None,
    header_row: int | None,
) -> Dataset:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if sheet_name:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(
                f"工作表“{sheet_name}”不存在。可用工作表："
                + ", ".join(workbook.sheetnames)
            )
        sheet = workbook[sheet_name]
    else:
        sheet = workbook.worksheets[0]
    if (
        sheet.max_row is None
        or sheet.max_column is None
        or (sheet.max_row <= 1 and sheet.max_column <= 1)
    ):
        sheet.reset_dimensions()
    raw_rows = list(sheet.iter_rows(values_only=True))
    if not raw_rows:
        raise ValueError(f"空工作表：{path} / {sheet.title}")
    header_index = detect_header_index(raw_rows, requested_keys, header_row)
    headers = make_unique_headers(raw_rows[header_index])
    rows: list[dict[str, Any]] = []
    for source_row, raw in enumerate(
        raw_rows[header_index + 1 :], start=header_index + 2
    ):
        values = [
            raw[index] if index < len(raw) else None
            for index in range(len(headers))
        ]
        if any(normalized_text(value) for value in values):
            rows.append({"source_row": source_row, "values": values})
    dataset = finalize_dataset(
        path=path,
        label=label,
        sheet_name=sheet.title,
        encoding=None,
        delimiter=None,
        header_row=header_index + 1,
        headers=headers,
        rows=rows,
        requested_keys=requested_keys,
        requested_secondary=requested_secondary,
        requested_date=requested_date,
    )
    workbook.close()
    return dataset


def read_dataset(
    path: Path,
    label: str,
    requested_keys: list[str],
    requested_secondary: list[str],
    requested_date: str | None,
    sheet_name: str | None,
    header_row: int | None,
) -> Dataset:
    suffix = path.suffix.casefold()
    if suffix in {".csv", ".tsv"}:
        return read_csv_dataset(
            path,
            label,
            requested_keys,
            requested_secondary,
            requested_date,
            header_row,
        )
    if suffix in {".xlsx", ".xlsm"}:
        return read_xlsx_dataset(
            path,
            label,
            requested_keys,
            requested_secondary,
            requested_date,
            sheet_name,
            header_row,
        )
    raise ValueError(
        f"不支持的文件类型：{path.suffix}。仅支持 CSV、TSV、XLSX、XLSM。"
    )


def is_identifier_header(header: str) -> bool:
    key = header_key(header)
    return (
        key.startswith("匹配键")
        or key.startswith("规范化键")
        or key in AUTO_ID_ALIASES
        or "id" in key
        or any(token in key for token in ("编号", "编码", "账号", "订单号"))
        or key.endswith("code")
    )


def is_date_header(header: str) -> bool:
    key = header_key(header)
    return (
        key in DATE_ALIASES
        or any(token in key for token in ("日期", "时间"))
        or key.endswith("date")
        or key.endswith("time")
        or key.endswith("at")
    )


def choose_auto_normalizer(left_header: str, right_header: str) -> str:
    if is_identifier_header(left_header) or is_identifier_header(right_header):
        return "id"
    if is_date_header(left_header) or is_date_header(right_header):
        return "date"
    return "text"


def resolve_normalizers(
    requested: list[str],
    left: Dataset,
    right: Dataset,
    secondary: bool,
) -> list[str]:
    left_fields = left.secondary_fields if secondary else left.key_fields
    right_fields = right.secondary_fields if secondary else right.key_fields
    count = len(left_fields)
    if count == 0:
        return []
    modes = requested or ["auto"]
    if len(modes) == 1:
        modes = modes * count
    if len(modes) != count:
        option = "--secondary-normalizer" if secondary else "--key-normalizer"
        raise ValueError(
            f"{option} 应提供1次或与字段对数量一致的{count}次，实际为{len(modes)}次。"
        )
    resolved: list[str] = []
    for index, mode in enumerate(modes):
        if mode == "auto":
            resolved.append(
                choose_auto_normalizer(
                    left.headers[left_fields[index]],
                    right.headers[right_fields[index]],
                )
            )
        else:
            resolved.append(mode)
    return resolved


def extract_link_id(value: Any, case_sensitive: bool) -> str:
    text = normalized_text(value, case_sensitive=True)
    for pattern in LINK_ID_PATTERNS:
        match = pattern.search(text)
        if match:
            token = match.group(1)
            return token if case_sensitive else token.casefold()
    return ""


def canonical_id(value: Any, case_sensitive: bool) -> tuple[str, bool]:
    if value is None or isinstance(value, bool):
        return "", False
    if isinstance(value, int):
        return str(value), abs(value) >= 10**15
    if isinstance(value, float):
        if not math.isfinite(value) or not value.is_integer():
            return "", False
        return format(value, ".0f"), abs(value) >= 10**15

    text = normalized_text(value, case_sensitive=True)
    while text.startswith("'"):
        text = text[1:].lstrip()
    compact = re.sub(r"\s+", "", text)
    if not compact:
        return "", False
    linked = extract_link_id(compact, case_sensitive)
    if linked:
        return linked, False
    if re.fullmatch(r"\+?\d+", compact):
        return compact.lstrip("+"), False
    decimal_match = re.fullmatch(r"\+?(\d+)\.0+", compact)
    if decimal_match:
        return decimal_match.group(1), False
    if re.fullmatch(r"[+-]?\d+(?:\.\d+)?[eE][+-]?\d+", compact):
        try:
            number = Decimal(compact)
            if number == number.to_integral_value() and number >= 0:
                return format(number.quantize(Decimal("1")), "f"), False
        except InvalidOperation:
            pass
    return (compact if case_sensitive else compact.casefold()), False


def normalize_number(value: Any) -> str:
    if value is None or isinstance(value, bool):
        return ""
    text = normalized_text(value, case_sensitive=True).replace(",", "")
    if not text:
        return ""
    try:
        number = Decimal(text)
    except InvalidOperation:
        return ""
    if not number.is_finite():
        return ""
    if number == 0:
        return "0"
    normalized = format(number.normalize(), "f")
    if "." in normalized:
        normalized = normalized.rstrip("0").rstrip(".")
    return normalized


def parse_datetime(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    text = normalized_text(value, case_sensitive=True)
    if not text:
        return None
    text = (
        text.replace("年", "-")
        .replace("月", "-")
        .replace("日", "")
        .replace("/", "-")
        .replace("T", " ")
        .rstrip("Z")
    )
    try:
        return datetime.fromisoformat(text).replace(tzinfo=None)
    except ValueError:
        match = re.fullmatch(
            r"(\d{4})-(\d{1,2})-(\d{1,2})(?:\s+(\d{1,2}):(\d{1,2})(?::(\d{1,2}))?)?",
            text,
        )
        if not match:
            return None
        parts = [int(value or 0) for value in match.groups()]
        try:
            return datetime(
                parts[0], parts[1], parts[2], parts[3], parts[4], parts[5]
            )
        except ValueError:
            return None


def normalize_value(
    value: Any, mode: str, case_sensitive: bool
) -> tuple[str, bool]:
    if mode == "exact":
        text = scalar_text(value)
        return (text if text != "" else ""), False
    if mode == "text":
        return normalized_text(value, case_sensitive), False
    if mode == "id":
        return canonical_id(value, case_sensitive)
    if mode == "number":
        return normalize_number(value), False
    if mode in {"date", "datetime"}:
        parsed = parse_datetime(value)
        if parsed is None:
            return "", False
        if mode == "date":
            return parsed.strftime("%Y-%m-%d"), False
        return parsed.strftime("%Y-%m-%d %H:%M:%S"), False
    raise ValueError(f"未知标准化模式：{mode}")


def prepare_records(
    dataset: Dataset,
    key_normalizers: list[str],
    secondary_normalizers: list[str],
    case_sensitive: bool,
) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    for item in dataset.rows:
        raw_key = tuple(
            scalar_text(item["values"][index]) for index in dataset.key_fields
        )
        trimmed_key = tuple(
            normalized_text(item["values"][index])
            for index in dataset.key_fields
        )
        normalized_components: list[str] = []
        precision_risks = 0
        for field_index, mode in zip(dataset.key_fields, key_normalizers):
            normalized, risk = normalize_value(
                item["values"][field_index], mode, case_sensitive
            )
            normalized_components.append(normalized)
            precision_risks += int(risk)

        secondary_components: list[str] = []
        for field_index, mode in zip(
            dataset.secondary_fields, secondary_normalizers
        ):
            normalized, _ = normalize_value(
                item["values"][field_index], mode, case_sensitive
            )
            secondary_components.append(normalized)

        key = tuple(normalized_components)
        secondary_key = tuple(secondary_components)
        records.append(
            {
                "source_row": item["source_row"],
                "values": item["values"],
                "raw_key": raw_key,
                "trimmed_key": trimmed_key,
                "key": key,
                "key_valid": bool(key) and all(key),
                "secondary_key": secondary_key,
                "secondary_valid": bool(secondary_key)
                and all(secondary_key),
                "publish_time": (
                    item["values"][dataset.date_field]
                    if dataset.date_field is not None
                    else None
                ),
                "precision_risks": precision_risks,
            }
        )
    return records


def group_by_key(
    records: list[dict[str, Any]]
) -> dict[tuple[str, ...], list[dict[str, Any]]]:
    grouped: dict[tuple[str, ...], list[dict[str, Any]]] = defaultdict(list)
    for record in records:
        if record["key_valid"]:
            grouped[record["key"]].append(record)
    return grouped


def valid_tuple_set(
    records: list[dict[str, Any]], field: str
) -> set[tuple[str, ...]]:
    return {
        record[field]
        for record in records
        if record[field] and all(record[field])
    }


def parse_time(value: Any) -> float | None:
    parsed = parse_datetime(value)
    return parsed.timestamp() if parsed else None


def display_key(key: tuple[str, ...]) -> str:
    return " | ".join(key)


def pair_secondary(
    left_records: list[dict[str, Any]],
    right_records: list[dict[str, Any]],
) -> tuple[
    list[dict[str, Any]],
    dict[int, dict[str, Any]],
    dict[int, dict[str, Any]],
    int,
]:
    left_groups: dict[tuple[str, ...], list[dict[str, Any]]] = defaultdict(list)
    right_groups: dict[tuple[str, ...], list[dict[str, Any]]] = defaultdict(list)
    for record in left_records:
        if record["secondary_valid"]:
            left_groups[record["secondary_key"]].append(record)
    for record in right_records:
        if record["secondary_valid"]:
            right_groups[record["secondary_key"]].append(record)

    common = set(left_groups) & set(right_groups)
    pairs: list[dict[str, Any]] = []
    left_map: dict[int, dict[str, Any]] = {}
    right_map: dict[int, dict[str, Any]] = {}
    for key in sorted(common):
        left = sorted(left_groups[key], key=lambda row: row["source_row"])
        remaining = sorted(
            right_groups[key], key=lambda row: row["source_row"]
        )
        for left_record in left:
            if not remaining:
                break
            left_time = parse_time(left_record["publish_time"])
            right_record = min(
                remaining,
                key=lambda row: (
                    left_time is None or parse_time(row["publish_time"]) is None,
                    abs(
                        (left_time or 0)
                        - (parse_time(row["publish_time"]) or 0)
                    ),
                    row["source_row"],
                ),
            )
            remaining.remove(right_record)
            pair = {
                "secondary_key": key,
                "left_source_row": left_record["source_row"],
                "right_source_row": right_record["source_row"],
                "left_key": left_record["key"],
                "right_key": right_record["key"],
            }
            pairs.append(pair)
            left_map[left_record["source_row"]] = pair
            right_map[right_record["source_row"]] = pair
    return pairs, left_map, right_map, len(common)


def duplicate_groups(records: list[dict[str, Any]]) -> int:
    return sum(1 for rows in group_by_key(records).values() if len(rows) > 1)


def date_range(records: list[dict[str, Any]]) -> tuple[str, str]:
    values = [
        parsed.strftime("%Y-%m-%d")
        for record in records
        if (parsed := parse_datetime(record["publish_time"])) is not None
    ]
    return (min(values), max(values)) if values else ("", "")


def analyze(
    left: Dataset,
    right: Dataset,
    left_records: list[dict[str, Any]],
    right_records: list[dict[str, Any]],
    key_specs: list[dict[str, str]],
    secondary_specs: list[dict[str, str]],
) -> dict[str, Any]:
    left_by_key = group_by_key(left_records)
    right_by_key = group_by_key(right_records)
    left_keys = set(left_by_key)
    right_keys = set(right_by_key)
    matched_keys = left_keys & right_keys
    left_only_keys = left_keys - right_keys
    right_only_keys = right_keys - left_keys

    matched_rows: list[list[Any]] = []
    for left_record in left_records:
        if not left_record["key_valid"]:
            continue
        for right_record in right_by_key.get(left_record["key"], []):
            matched_rows.append(
                list(left_record["key"])
                + [left_record["source_row"], right_record["source_row"]]
                + left_record["values"]
                + right_record["values"]
            )

    left_unmatched = [
        record
        for record in left_records
        if not record["key_valid"] or record["key"] in left_only_keys
    ]
    right_unmatched = [
        record
        for record in right_records
        if not record["key_valid"] or record["key"] in right_only_keys
    ]
    if secondary_specs:
        (
            secondary_pairs,
            left_pair_map,
            right_pair_map,
            secondary_common_groups,
        ) = pair_secondary(left_unmatched, right_unmatched)
    else:
        secondary_pairs = []
        left_pair_map = {}
        right_pair_map = {}
        secondary_common_groups = 0

    matched_key_headers = [
        f"匹配键{index}_{spec['left']}={spec['right']}"
        for index, spec in enumerate(key_specs, start=1)
    ]
    matched_headers = (
        matched_key_headers
        + ["文件A源行号", "文件B源行号"]
        + [f"文件A_{header}" for header in left.headers]
        + [f"文件B_{header}" for header in right.headers]
    )
    left_unmatched_key_headers = [
        f"规范化键{index}_{spec['left']}"
        for index, spec in enumerate(key_specs, start=1)
    ]
    right_unmatched_key_headers = [
        f"规范化键{index}_{spec['right']}"
        for index, spec in enumerate(key_specs, start=1)
    ]
    left_unmatched_prefix = (
        ["来源行号", "匹配状态"]
        + left_unmatched_key_headers
        + [
            "次级字段补充匹配",
            "补充匹配对方行号",
            "补充匹配对方主键",
        ]
    )
    right_unmatched_prefix = (
        ["来源行号", "匹配状态"]
        + right_unmatched_key_headers
        + [
            "次级字段补充匹配",
            "补充匹配对方行号",
            "补充匹配对方主键",
        ]
    )

    left_unmatched_rows: list[list[Any]] = []
    for record in left_unmatched:
        pair = left_pair_map.get(record["source_row"])
        left_unmatched_rows.append(
            [
                record["source_row"],
                "仅文件A存在" if record["key_valid"] else "匹配键为空或格式异常",
            ]
            + list(record["key"])
            + [
                "是" if pair else ("否" if secondary_specs else "未启用"),
                pair["right_source_row"] if pair else None,
                display_key(pair["right_key"]) if pair else None,
            ]
            + record["values"]
        )
    right_unmatched_rows: list[list[Any]] = []
    for record in right_unmatched:
        pair = right_pair_map.get(record["source_row"])
        right_unmatched_rows.append(
            [
                record["source_row"],
                "仅文件B存在" if record["key_valid"] else "匹配键为空或格式异常",
            ]
            + list(record["key"])
            + [
                "是" if pair else ("否" if secondary_specs else "未启用"),
                pair["left_source_row"] if pair else None,
                display_key(pair["left_key"]) if pair else None,
            ]
            + record["values"]
        )

    if left_keys == right_keys:
        relation = "两个文件的有效规范化匹配键集合完全一致"
    elif right_keys < left_keys:
        relation = "文件B是文件A的严格子集"
    elif left_keys < right_keys:
        relation = "文件A是文件B的严格子集"
    elif matched_keys:
        relation = "两个文件部分重叠，互不包含"
    else:
        relation = "两个文件不存在共同有效规范化匹配键"

    summary = {
        "left_records": len(left_records),
        "right_records": len(right_records),
        "left_valid_key_records": sum(
            record["key_valid"] for record in left_records
        ),
        "right_valid_key_records": sum(
            record["key_valid"] for record in right_records
        ),
        "left_unique_keys": len(left_keys),
        "right_unique_keys": len(right_keys),
        "matched_unique_keys": len(matched_keys),
        "inner_join_rows": len(matched_rows),
        "left_unmatched_records": len(left_unmatched),
        "right_unmatched_records": len(right_unmatched),
        "left_only_unique_keys": len(left_only_keys),
        "right_only_unique_keys": len(right_only_keys),
        "left_invalid_key_records": sum(
            not record["key_valid"] for record in left_records
        ),
        "right_invalid_key_records": sum(
            not record["key_valid"] for record in right_records
        ),
        "left_duplicate_key_groups": duplicate_groups(left_records),
        "right_duplicate_key_groups": duplicate_groups(right_records),
        "left_match_rate": len(matched_keys) / len(left_keys) if left_keys else 0,
        "right_match_rate": len(matched_keys) / len(right_keys)
        if right_keys
        else 0,
        "left_is_subset_of_right": left_keys <= right_keys,
        "right_is_subset_of_left": right_keys <= left_keys,
        "relation": relation,
        "secondary_common_groups": secondary_common_groups,
        "secondary_additional_pairs": len(secondary_pairs),
        "adjusted_left_unmatched_records": len(left_unmatched)
        - len(secondary_pairs),
        "adjusted_right_unmatched_records": len(right_unmatched)
        - len(secondary_pairs),
    }
    audit = {
        "raw_exact_key_intersection": len(
            valid_tuple_set(left_records, "raw_key")
            & valid_tuple_set(right_records, "raw_key")
        ),
        "trimmed_key_intersection": len(
            valid_tuple_set(left_records, "trimmed_key")
            & valid_tuple_set(right_records, "trimmed_key")
        ),
        "normalized_key_intersection": len(matched_keys),
        "left_changed_key_records": sum(
            record["key_valid"] and record["raw_key"] != record["key"]
            for record in left_records
        ),
        "right_changed_key_records": sum(
            record["key_valid"] and record["raw_key"] != record["key"]
            for record in right_records
        ),
        "left_numeric_precision_risks": sum(
            record["precision_risks"] for record in left_records
        ),
        "right_numeric_precision_risks": sum(
            record["precision_risks"] for record in right_records
        ),
    }
    key_count = len(key_specs)
    supplement_column = 3 + key_count
    return {
        "summary": summary,
        "audit": audit,
        "matched_keys": matched_keys,
        "left_keys": left_keys,
        "right_keys": right_keys,
        "key_specs": key_specs,
        "secondary_specs": secondary_specs,
        "date_ranges": {
            "left": date_range(left_records),
            "right": date_range(right_records),
        },
        "sheets": [
            {
                "name": "匹配表_inner join",
                "headers": matched_headers,
                "rows": matched_rows,
                "theme": "matched",
                "freeze_columns": key_count + 2,
                "key_columns": list(range(1, key_count + 1)),
                "supplement_column": None,
            },
            {
                "name": "不匹配详情_文件A",
                "headers": left_unmatched_prefix + left.headers,
                "rows": left_unmatched_rows,
                "theme": "left",
                "freeze_columns": key_count + 5,
                "key_columns": list(range(3, 3 + key_count)),
                "supplement_column": supplement_column,
            },
            {
                "name": "不匹配详情_文件B",
                "headers": right_unmatched_prefix + right.headers,
                "rows": right_unmatched_rows,
                "theme": "right",
                "freeze_columns": key_count + 5,
                "key_columns": list(range(3, 3 + key_count)),
                "supplement_column": supplement_column,
            },
        ],
    }


def safe_excel_value(header: str, value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, datetime) and value.tzinfo is not None:
        value = value.replace(tzinfo=None)
    if isinstance(value, (datetime, date, int, float, bool)):
        if not is_identifier_header(header):
            return value
    if isinstance(value, (list, dict, tuple, set)):
        value = json.dumps(value, ensure_ascii=False)
    text = ILLEGAL_CHARACTERS_RE.sub("", str(value))
    return text[:32767]


def display_width(value: Any) -> int:
    text = scalar_text(value)
    return sum(
        2 if unicodedata.east_asian_width(char) in {"W", "F"} else 1
        for char in text
    )


def width_for_column(header: str, values: list[Any]) -> float:
    if is_identifier_header(header):
        return 28
    key = header_key(header)
    if "链接" in header or key in {"url", "link"}:
        return 38
    if "标题" in header:
        return 34
    if any(token in header for token in ("正文", "摘要", "说明")):
        return 40
    if is_date_header(header):
        return 20
    sample_width = max(
        [display_width(header)]
        + [display_width(value) for value in values[:200]]
    )
    return min(max(sample_width + 2, 10), 24)


def append_typed_row(ws, headers: list[str], values: list[Any]) -> None:
    prepared = [
        safe_excel_value(headers[index], value)
        for index, value in enumerate(values)
    ]
    ws.append(prepared)
    for cell in ws[ws.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = "s"


def style_data_sheet(
    ws,
    sheet_data: dict[str, Any],
    table_index: int,
) -> None:
    headers = sheet_data["headers"]
    rows = sheet_data["rows"]
    theme = sheet_data["theme"]
    colors = {
        "matched": ("2E7D32", "E8F5E9"),
        "left": ("9F1239", "FFF1F2"),
        "right": ("B45309", "FFF7ED"),
    }
    header_color, light_color = colors[theme]
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = (
        f"{get_column_letter(sheet_data['freeze_columns'] + 1)}2"
    )
    ws.auto_filter.ref = (
        f"A1:{get_column_letter(len(headers))}{max(len(rows) + 1, 1)}"
    )
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=header_color)
        cell.font = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
    ws.row_dimensions[1].height = 38

    key_columns = set(sheet_data["key_columns"])
    for index, header in enumerate(headers, start=1):
        values = [row[index - 1] for row in rows if index <= len(row)]
        ws.column_dimensions[get_column_letter(index)].width = width_for_column(
            header, values
        )
        if rows and (index in key_columns or is_identifier_header(header)):
            for column_cells in ws.iter_cols(
                min_col=index,
                max_col=index,
                min_row=2,
                max_row=len(rows) + 1,
            ):
                for cell in column_cells:
                    cell.number_format = "0"

    if rows:
        ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
        table = Table(displayName=f"RecordMatch{table_index}", ref=ref)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(table)

    if theme in {"left", "right"} and rows:
        supplement_column = sheet_data["supplement_column"]
        for row_index in range(2, len(rows) + 2):
            ws.cell(row_index, 2).fill = PatternFill(
                "solid", fgColor=light_color
            )
            ws.cell(row_index, 2).font = Font(
                name="Aptos", size=10, bold=True, color=header_color
            )
            if (
                supplement_column is not None
                and ws.cell(row_index, supplement_column).value == "是"
            ):
                for column in range(
                    supplement_column, supplement_column + 3
                ):
                    ws.cell(row_index, column).fill = PatternFill(
                        "solid", fgColor="E0F2FE"
                    )
                    ws.cell(row_index, column).font = Font(
                        name="Aptos", size=10, bold=True, color="075985"
                    )


def add_section(ws, row: int, title: str) -> None:
    ws.merge_cells(
        start_row=row, start_column=1, end_row=row, end_column=4
    )
    cell = ws.cell(row, 1, title)
    cell.fill = PatternFill("solid", fgColor="1F4E78")
    cell.font = Font(name="Aptos", size=11, bold=True, color="FFFFFF")
    cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[row].height = 26


def add_analysis_sheet(
    workbook: Workbook,
    left: Dataset,
    right: Dataset,
    result: dict[str, Any],
) -> None:
    ws = workbook.create_sheet("分析汇总")
    ws.sheet_view.showGridLines = False
    summary = result["summary"]
    audit = result["audit"]
    left_start, left_end = result["date_ranges"]["left"]
    right_start, right_end = result["date_ranges"]["right"]

    ws.merge_cells("A1:H1")
    ws["A1"] = "字段匹配与差异分析汇总"
    ws["A1"].fill = PatternFill("solid", fgColor="17365D")
    ws["A1"].font = Font(
        name="Aptos Display", size=18, bold=True, color="FFFFFF"
    )
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 38
    ws.merge_cells("A2:H2")
    ws["A2"] = (
        f"文件A：{left.path}    文件B：{right.path}    "
        f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    )
    ws["A2"].font = Font(name="Aptos", size=9, color="475569")
    ws["A2"].alignment = Alignment(vertical="center")
    ws.row_dimensions[2].height = 24

    rows: list[tuple[Any, Any, Any, Any]] = []
    rows.append(("主匹配配置", None, None, None))
    rows.append(("指标", left.label, right.label, "标准化方式/说明"))
    for index, spec in enumerate(result["key_specs"], start=1):
        rows.append(
            (
                f"主键组件{index}",
                spec["left"],
                spec["right"],
                spec["normalizer"],
            )
        )
    if result["secondary_specs"]:
        for index, spec in enumerate(result["secondary_specs"], start=1):
            rows.append(
                (
                    f"次级键组件{index}",
                    spec["left"],
                    spec["right"],
                    spec["normalizer"],
                )
            )
    rows.append(
        (
            "日期字段",
            left.headers[left.date_field] if left.date_field is not None else "",
            right.headers[right.date_field]
            if right.date_field is not None
            else "",
            "用于日期范围与次级配对顺序",
        )
    )
    rows.extend(
        [
            ("匹配概览", None, None, None),
            ("指标", left.label, right.label, "说明"),
            (
                "源数据记录数",
                summary["left_records"],
                summary["right_records"],
                "不含表头",
            ),
            (
                "有效主键记录数",
                summary["left_valid_key_records"],
                summary["right_valid_key_records"],
                "所有主键组件均有效",
            ),
            (
                "唯一有效主键数",
                summary["left_unique_keys"],
                summary["right_unique_keys"],
                "复合键按元组去重",
            ),
            (
                "匹配唯一主键数",
                summary["matched_unique_keys"],
                summary["matched_unique_keys"],
                "规范化主键交集",
            ),
            (
                "inner join行数",
                summary["inner_join_rows"],
                summary["inner_join_rows"],
                "重复键采用笛卡尔连接",
            ),
            (
                "不匹配记录数",
                summary["left_unmatched_records"],
                summary["right_unmatched_records"],
                "包含无效主键",
            ),
            (
                "仅本文件唯一主键数",
                summary["left_only_unique_keys"],
                summary["right_only_unique_keys"],
                "方向性差异",
            ),
            (
                "匹配率/覆盖率",
                summary["left_match_rate"],
                summary["right_match_rate"],
                "匹配唯一主键 / 本文件唯一有效主键",
            ),
            (
                "重复主键组数",
                summary["left_duplicate_key_groups"],
                summary["right_duplicate_key_groups"],
                "同一规范化主键出现多次",
            ),
            (
                "无效主键记录数",
                summary["left_invalid_key_records"],
                summary["right_invalid_key_records"],
                "任一主键组件为空或无效",
            ),
            (
                "有效主键集合包含关系",
                "是" if summary["left_is_subset_of_right"] else "否",
                "是" if summary["right_is_subset_of_left"] else "否",
                "本文件是否为对方子集",
            ),
            (
                "日期范围",
                f"{left_start} 至 {left_end}" if left_start else "",
                f"{right_start} 至 {right_end}" if right_start else "",
                "自动或显式日期字段",
            ),
            ("标准化审计", None, None, None),
            ("指标", left.label, right.label, "说明"),
            (
                "原始主键集合交集",
                audit["raw_exact_key_intersection"],
                audit["raw_exact_key_intersection"],
                "源值转字符串后比较",
            ),
            (
                "文本清洗主键集合交集",
                audit["trimmed_key_intersection"],
                audit["trimmed_key_intersection"],
                "NFKC、空白折叠、大小写归一",
            ),
            (
                "最终规范化主键集合交集",
                audit["normalized_key_intersection"],
                audit["normalized_key_intersection"],
                "最终精确匹配口径",
            ),
            (
                "标准化后发生变化的记录",
                audit["left_changed_key_records"],
                audit["right_changed_key_records"],
                "原始主键与最终主键不同",
            ),
            (
                "Excel长数字精度风险",
                audit["left_numeric_precision_risks"],
                audit["right_numeric_precision_risks"],
                "ID模式下数值型16位及以上整数",
            ),
            ("次级匹配与结论", None, None, None),
            ("指标", left.label, right.label, "说明"),
            (
                "次级键共同组数",
                summary["secondary_common_groups"],
                summary["secondary_common_groups"],
                "仅主键不匹配记录",
            ),
            (
                "次级字段补充匹配对",
                summary["secondary_additional_pairs"],
                summary["secondary_additional_pairs"],
                "一对一配对，不计入inner join",
            ),
            (
                "补充判断后仍不匹配",
                summary["adjusted_left_unmatched_records"],
                summary["adjusted_right_unmatched_records"],
                "不改变主键精确匹配口径",
            ),
            (
                "最终结论",
                summary["relation"],
                summary["relation"],
                "按唯一有效规范化主键集合判断",
            ),
        ]
    )

    start_row = 4
    for offset, row in enumerate(rows):
        row_index = start_row + offset
        if row[1:] == (None, None, None):
            add_section(ws, row_index, row[0])
            continue
        for column, value in enumerate(row, start=1):
            cell = ws.cell(row_index, column, value)
            cell.font = Font(name="Aptos", size=10, color="1F2937")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = Border(
                bottom=Side(style="thin", color="E2E8F0")
            )
        if row[0] == "指标":
            for column in range(1, 5):
                cell = ws.cell(row_index, column)
                cell.fill = PatternFill("solid", fgColor="334155")
                cell.font = Font(
                    name="Aptos", size=10, bold=True, color="FFFFFF"
                )
                cell.alignment = Alignment(
                    horizontal="center", vertical="center", wrap_text=True
                )
            ws.row_dimensions[row_index].height = 28
        elif row[0] == "匹配率/覆盖率":
            ws.cell(row_index, 2).number_format = "0.00%"
            ws.cell(row_index, 3).number_format = "0.00%"

    ws.freeze_panes = "A4"
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 34
    ws.column_dimensions["D"].width = 40
    for column in "EFGH":
        ws.column_dimensions[column].width = 4


def write_workbook(
    output: Path,
    left: Dataset,
    right: Dataset,
    result: dict[str, Any],
) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for table_index, sheet_data in enumerate(result["sheets"], start=1):
        ws = workbook.create_sheet(sheet_data["name"])
        append_typed_row(ws, sheet_data["headers"], sheet_data["headers"])
        for row in sheet_data["rows"]:
            append_typed_row(ws, sheet_data["headers"], row)
        style_data_sheet(ws, sheet_data, table_index)
    add_analysis_sheet(workbook, left, right, result)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def validate_output(output: Path, result: dict[str, Any]) -> dict[str, Any]:
    with zipfile.ZipFile(output) as archive:
        bad_member = archive.testzip()
        if bad_member is not None:
            raise AssertionError(f"OOXML压缩包损坏：{bad_member}")

    workbook = load_workbook(output, read_only=False, data_only=False)
    expected_names = [
        "匹配表_inner join",
        "不匹配详情_文件A",
        "不匹配详情_文件B",
        "分析汇总",
    ]
    if workbook.sheetnames != expected_names:
        raise AssertionError(f"工作表顺序错误：{workbook.sheetnames}")
    if any(workbook[name].sheet_state != "visible" for name in expected_names):
        raise AssertionError("存在非可见输出工作表")

    for sheet_data in result["sheets"]:
        ws = workbook[sheet_data["name"]]
        expected_rows = len(sheet_data["rows"]) + 1
        expected_columns = len(sheet_data["headers"])
        if ws.max_row != expected_rows:
            raise AssertionError(
                f"{ws.title} 行数错误：{ws.max_row} != {expected_rows}"
            )
        if ws.max_column != expected_columns:
            raise AssertionError(
                f"{ws.title} 列数错误：{ws.max_column} != {expected_columns}"
            )

    key_count = len(result["key_specs"])
    matched_ws = workbook["匹配表_inner join"]
    matched_keys = {
        tuple(
            matched_ws.cell(row, column).value
            for column in range(1, key_count + 1)
        )
        for row in range(2, matched_ws.max_row + 1)
    }
    if matched_keys != result["matched_keys"]:
        raise AssertionError("匹配表主键集合与计算结果不一致")

    unmatched_checks = [
        ("不匹配详情_文件A", result["left_keys"] - result["matched_keys"]),
        ("不匹配详情_文件B", result["right_keys"] - result["matched_keys"]),
    ]
    for sheet_name, expected_keys in unmatched_checks:
        ws = workbook[sheet_name]
        actual_keys = set()
        for row in range(2, ws.max_row + 1):
            key = tuple(
                ws.cell(row, column).value
                for column in range(3, 3 + key_count)
            )
            if all(key):
                actual_keys.add(key)
        if actual_keys != expected_keys:
            raise AssertionError(f"{sheet_name} 主键集合与计算结果不一致")

    for sheet_data in result["sheets"]:
        ws = workbook[sheet_data["name"]]
        for column in sheet_data["key_columns"]:
            for row in range(2, ws.max_row + 1):
                value = ws.cell(row, column).value
                if value is not None and not isinstance(value, str):
                    raise AssertionError(
                        f"规范化主键未按文本存储："
                        f"{ws.title}!{get_column_letter(column)}{row}"
                    )

    for ws in workbook.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "e":
                    raise AssertionError(
                        f"发现Excel错误值：{ws.title}!{cell.coordinate}"
                    )
                if isinstance(cell.value, str) and cell.value.startswith(
                    EXCEL_ERROR_PREFIXES
                ):
                    raise AssertionError(
                        f"发现错误文本：{ws.title}!{cell.coordinate}"
                    )
    return {
        "zip_integrity": "passed",
        "sheet_order": "passed",
        "shape_checks": "passed",
        "key_set_checks": "passed",
        "key_text_checks": "passed",
        "error_scan": "passed",
    }


def build_specs(
    left: Dataset,
    right: Dataset,
    normalizers: list[str],
    secondary: bool,
) -> list[dict[str, str]]:
    left_fields = left.secondary_fields if secondary else left.key_fields
    right_fields = right.secondary_fields if secondary else right.key_fields
    return [
        {
            "left": left.headers[left_index],
            "right": right.headers[right_index],
            "normalizer": normalizer,
        }
        for left_index, right_index, normalizer in zip(
            left_fields, right_fields, normalizers
        )
    ]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "按一个字段或多个字段组成的复合键匹配两个CSV/XLSX文件，"
            "输出inner join、双向差异、次级判断和分析汇总。"
        )
    )
    parser.add_argument("left_file", type=Path, help="文件A路径")
    parser.add_argument("right_file", type=Path, help="文件B路径")
    parser.add_argument("--output", type=Path, help="输出XLSX路径")
    parser.add_argument("--json-report", type=Path, help="可选JSON结果路径")
    parser.add_argument("--left-sheet", help="文件A工作表名称")
    parser.add_argument("--right-sheet", help="文件B工作表名称")
    parser.add_argument("--left-header-row", type=int, help="文件A表头行，1起始")
    parser.add_argument("--right-header-row", type=int, help="文件B表头行，1起始")
    parser.add_argument("--left-label", default="文件A", help="文件A显示名称")
    parser.add_argument("--right-label", default="文件B", help="文件B显示名称")
    parser.add_argument(
        "--left-key",
        action="append",
        default=[],
        help="文件A主匹配字段；复合键可重复提供",
    )
    parser.add_argument(
        "--right-key",
        action="append",
        default=[],
        help="文件B主匹配字段；须与--left-key按位置对应",
    )
    parser.add_argument(
        "--key-normalizer",
        action="append",
        choices=NORMALIZER_CHOICES,
        default=[],
        help="主匹配字段标准化方式；可提供1次或每个字段对1次",
    )
    parser.add_argument(
        "--left-secondary-key",
        action="append",
        default=[],
        help="文件A次级匹配字段；可重复提供",
    )
    parser.add_argument(
        "--right-secondary-key",
        action="append",
        default=[],
        help="文件B次级匹配字段；须按位置对应",
    )
    parser.add_argument(
        "--secondary-normalizer",
        action="append",
        choices=NORMALIZER_CHOICES,
        default=[],
        help="次级匹配字段标准化方式",
    )
    parser.add_argument("--left-date", help="文件A日期/时间字段")
    parser.add_argument("--right-date", help="文件B日期/时间字段")
    parser.add_argument(
        "--case-sensitive",
        action="store_true",
        help="text和id模式区分大小写；默认不区分",
    )
    parser.add_argument(
        "--left-id",
        help=argparse.SUPPRESS,
    )
    parser.add_argument(
        "--right-id",
        help=argparse.SUPPRESS,
    )
    parser.add_argument(
        "--id-case-sensitive",
        dest="case_sensitive",
        action="store_true",
        help=argparse.SUPPRESS,
    )
    return parser.parse_args()


def validate_cli_mappings(args: argparse.Namespace) -> tuple[
    list[str], list[str], list[str], list[str]
]:
    left_keys = list(args.left_key)
    right_keys = list(args.right_key)
    if not left_keys and args.left_id:
        left_keys = [args.left_id]
    if not right_keys and args.right_id:
        right_keys = [args.right_id]
    if bool(left_keys) != bool(right_keys):
        raise ValueError("文件A和文件B必须同时提供主匹配字段。")
    if left_keys and len(left_keys) != len(right_keys):
        raise ValueError(
            "--left-key 与 --right-key 数量必须一致且按位置对应。"
        )

    left_secondary = list(args.left_secondary_key)
    right_secondary = list(args.right_secondary_key)
    if bool(left_secondary) != bool(right_secondary):
        raise ValueError("文件A和文件B必须同时提供次级匹配字段。")
    if len(left_secondary) != len(right_secondary):
        raise ValueError(
            "--left-secondary-key 与 --right-secondary-key 数量必须一致。"
        )
    return left_keys, right_keys, left_secondary, right_secondary


def main() -> int:
    args = parse_args()
    (
        left_keys,
        right_keys,
        left_secondary,
        right_secondary,
    ) = validate_cli_mappings(args)
    left_path = args.left_file.resolve()
    right_path = args.right_file.resolve()
    if not left_path.is_file():
        raise FileNotFoundError(f"文件A不存在：{left_path}")
    if not right_path.is_file():
        raise FileNotFoundError(f"文件B不存在：{right_path}")
    output = (
        args.output.resolve()
        if args.output
        else (
            Path.cwd()
            / f"{left_path.stem}_{right_path.stem}_字段匹配结果.xlsx"
        ).resolve()
    )
    if output in {left_path, right_path}:
        raise ValueError("输出路径不得覆盖源文件")

    left = read_dataset(
        left_path,
        args.left_label,
        left_keys,
        left_secondary,
        args.left_date,
        args.left_sheet,
        args.left_header_row,
    )
    right = read_dataset(
        right_path,
        args.right_label,
        right_keys,
        right_secondary,
        args.right_date,
        args.right_sheet,
        args.right_header_row,
    )
    if len(left.key_fields) != len(right.key_fields):
        raise ValueError("自动识别后的主匹配字段数量不一致。")
    if len(left.secondary_fields) != len(right.secondary_fields):
        raise ValueError("次级匹配字段数量不一致。")

    key_normalizers = resolve_normalizers(
        args.key_normalizer, left, right, secondary=False
    )
    secondary_normalizers = resolve_normalizers(
        args.secondary_normalizer, left, right, secondary=True
    )
    key_specs = build_specs(
        left, right, key_normalizers, secondary=False
    )
    secondary_specs = build_specs(
        left, right, secondary_normalizers, secondary=True
    )
    left_records = prepare_records(
        left,
        key_normalizers,
        secondary_normalizers,
        args.case_sensitive,
    )
    right_records = prepare_records(
        right,
        key_normalizers,
        secondary_normalizers,
        args.case_sensitive,
    )
    result = analyze(
        left,
        right,
        left_records,
        right_records,
        key_specs,
        secondary_specs,
    )
    write_workbook(output, left, right, result)
    validation = validate_output(output, result)

    report = {
        "output": str(output),
        "sources": {
            "left": str(left.path),
            "right": str(right.path),
        },
        "source_metadata": {
            "left_sheet": left.sheet_name,
            "right_sheet": right.sheet_name,
            "left_header_row": left.header_row,
            "right_header_row": right.header_row,
            "left_encoding": left.encoding,
            "right_encoding": right.encoding,
        },
        "field_mapping": {
            "primary_keys": key_specs,
            "secondary_keys": secondary_specs,
            "left_date": (
                left.headers[left.date_field]
                if left.date_field is not None
                else None
            ),
            "right_date": (
                right.headers[right.date_field]
                if right.date_field is not None
                else None
            ),
        },
        "summary": result["summary"],
        "normalization_audit": result["audit"],
        "validation": validation,
    }
    if args.json_report:
        json_path = args.json_report.resolve()
        json_path.parent.mkdir(parents=True, exist_ok=True)
        json_path.write_text(
            json.dumps(report, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
